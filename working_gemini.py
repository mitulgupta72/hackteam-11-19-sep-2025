import os
import pytesseract
from PIL import Image
import PyPDF2
import docx
import pdfplumber
import openpyxl  # For handling .xlsx files
import xlrd  # For handling .xls files
import csv  # For handling CSV files
from langchain.text_splitter import RecursiveCharacterTextSplitter
import requests
import json
from pinecone import Pinecone, ServerlessSpec

# Set path for Tesseract OCR
pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'

import time

# Set up your API Key from Google Cloud (Gemini API Key)
API_KEY = "AIzaSyBgJglaZXlsG7JTRgGCmvY18yS07J_Q7Ng"  # Replace with your Gemini (PaLM) API key
ENDPOINT = "https://generativelanguage.googleapis.com/v1beta/models/embedding-001:embedContent"  # Fixed Gemini endpoint

# Alternative OpenAI API key (if switching providers)
OPENAI_API_KEY = "your-openai-api-key-here"  # Get from https://platform.openai.com/api-keys

# Initialize Pinecone with your API key (FIXED)
pc = Pinecone(api_key="pcsk_5DD3n_EcYuCi2mpsqpD9RehYSMVGUApUndoQDjCFQgnwXRz539LCjj33xjbH2ckntz3H3")

# Check if the index exists; if not, create a new one
index_name = "audit19092025"
try:
    existing_indexes = [index.name for index in pc.list_indexes()]
    if index_name not in existing_indexes:
        pc.create_index(
            name=index_name,
            dimension=768,  # Gemini embedding dimension is 768, not 3072
            metric='cosine',
            spec=ServerlessSpec(cloud='aws', region='us-east-1')
        )
        print(f"Created new index: {index_name}")
    else:
        print(f"Index {index_name} already exists")
        
    # Connect to the index
    index = pc.Index(index_name)
    print(f"Successfully connected to index: {index_name}")
    
except Exception as e:
    print(f"Error with Pinecone setup: {e}")

# PDF to text extraction
def pdf_to_text(pdf_path):
    try:
        text = ""
        with open(pdf_path, "rb") as file:
            pdf_reader = PyPDF2.PdfReader(file)
            num_pages = len(pdf_reader.pages)
            for page_num in range(num_pages):
                page = pdf_reader.pages[page_num]
                text += page.extract_text()
        return text
    except Exception as e:
        return f"Error reading PDF: {e}"

# Image to text extraction using Tesseract OCR
def image_to_text(image_path):
    try:
        text = pytesseract.image_to_string(Image.open(image_path))
        return text
    except Exception as e:
        return f"Error reading image: {e}"

# DOCX to text extraction
def docs_to_text(docs_path):
    try:
        doc = docx.Document(docs_path)
        text = ""
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        return text
    except Exception as e:
        return f"Error reading DOCX: {e}"

# TXT to text extraction
def txt_to_text(txt_path):
    try:
        with open(txt_path, 'r', encoding='utf-8') as file:
            return file.read()
    except Exception as e:
        return f"Error reading TXT file: {e}"

# Excel (XLSX) to text extraction (using openpyxl)
def excel_to_text(excel_path):
    try:
        wb = openpyxl.load_workbook(excel_path)
        sheet = wb.active
        text = ""
        for row in sheet.iter_rows(values_only=True):
            text += "\t".join([str(cell) if cell is not None else "" for cell in row]) + "\n"
        return text
    except Exception as e:
        return f"Error reading Excel file (.xlsx): {e}"

# Excel (XLS) to text extraction (using xlrd)
def xls_to_text(xls_path):
    try:
        wb = xlrd.open_workbook(xls_path)
        sheet = wb.sheet_by_index(0)
        text = ""
        for row_num in range(sheet.nrows):
            row = sheet.row(row_num)
            text += "\t".join([str(cell.value) for cell in row]) + "\n"
        return text
    except Exception as e:
        return f"Error reading Excel file (.xls): {e}"

# CSV to text extraction
def csv_to_text(csv_path):
    try:
        with open(csv_path, newline='', encoding='utf-8') as file:
            reader = csv.reader(file)
            text = ""
            for row in reader:
                text += "\t".join(row) + "\n"
        return text
    except Exception as e:
        return f"Error reading CSV file: {e}"

# Main function to extract text based on file extension
def extract_text(file_path):
    _, ext = os.path.splitext(file_path)
    ext = ext.lower()

    if ext == ".pdf":
        return pdf_to_text(file_path)
    elif ext in (".jpg", ".jpeg", ".png", ".bmp", ".gif"):
        return image_to_text(file_path)
    elif ext == ".docx":
        return docs_to_text(file_path)
    elif ext == ".txt":
        return txt_to_text(file_path)
    elif ext == ".xlsx":
        return excel_to_text(file_path)
    elif ext == ".xls":
        return xls_to_text(file_path)
    elif ext == ".csv":
        return csv_to_text(file_path)
    else:
        return "Unsupported file type"

# Function to chunk the extracted text
def chunk_text(text):
    # Initialize the text splitter (you can adjust chunk size and overlap)
    text_splitter = RecursiveCharacterTextSplitter(chunk_size=500, chunk_overlap=100)
    chunks = text_splitter.split_text(text)
    return chunks

# Alternative embedding function using OpenAI (when Gemini quota is exceeded)
def generate_embeddings_openai(chunks):
    """
    Alternative embedding function using OpenAI API
    Install: pip install openai
    Get API key from: https://platform.openai.com/api-keys
    """
    try:
        import openai
        
        # Set your OpenAI API key
        openai.api_key = "your-openai-api-key-here"  # Replace with your actual key
        
        embeddings = []
        for i, chunk in enumerate(chunks):
            try:
                response = openai.Embedding.create(
                    model="text-embedding-ada-002",
                    input=chunk
                )
                embedding = response['data'][0]['embedding']
                embeddings.append(embedding)
                print(f"Generated OpenAI embedding for chunk {i+1}/{len(chunks)}")
            except Exception as e:
                print(f"Error generating OpenAI embedding for chunk {i+1}: {e}")
                embeddings.append(None)
        
        valid_embeddings = [(i, emb) for i, emb in enumerate(embeddings) if emb is not None]
        return valid_embeddings
        
    except ImportError:
        print("OpenAI library not installed. Run: pip install openai")
        return []

# Function to generate embeddings using the Gemini API with retry logic
def generate_embeddings_gemini(chunks, max_retries=3, delay=60):
    embeddings = []
    
    for i, chunk in enumerate(chunks):
        retries = 0
        success = False
        
        while retries < max_retries and not success:
            try:
                # Create the payload for the Gemini API request (FIXED format)
                payload = {
                    "model": "models/embedding-001",
                    "content": {
                        "parts": [{"text": chunk}]
                    }
                }
                
                headers = {
                    "x-goog-api-key": API_KEY,  # Fixed header format for Gemini
                    "Content-Type": "application/json"
                }
                
                # Make the API request to generate embeddings
                response = requests.post(ENDPOINT, headers=headers, data=json.dumps(payload))
                
                if response.status_code == 200:
                    result = response.json()
                    # Extract embedding from Gemini response format
                    embedding = result['embedding']['values']
                    embeddings.append(embedding)
                    print(f"Generated embedding for chunk {i+1}/{len(chunks)}")
                    success = True
                elif response.status_code == 429:  # Quota exceeded
                    print(f"Quota exceeded for chunk {i+1}. Retrying in {delay} seconds... (Retry {retries+1}/{max_retries})")
                    time.sleep(delay)
                    retries += 1
                    if retries >= max_retries:
                        print(f"Max retries reached for chunk {i+1}. Skipping.")
                        embeddings.append(None)
                else:
                    print(f"Error for chunk {i+1}: {response.status_code}, {response.text}")
                    embeddings.append(None)
                    break
                    
            except Exception as e:
                print(f"Exception for chunk {i+1}: {e}")
                embeddings.append(None)
                break
    
    # Filter out None values
    valid_embeddings = [(i, emb) for i, emb in enumerate(embeddings) if emb is not None]
    return valid_embeddings

# Function to push embeddings into Pinecone (FIXED)
def push_embeddings_to_pinecone(embeddings_with_indices, chunks):
    try:
        vectors_to_upsert = []
        for chunk_idx, embedding in embeddings_with_indices:
            vectors_to_upsert.append({
                'id': f"doc_{chunk_idx}",
                'values': embedding,
                'metadata': {
                    'text': chunks[chunk_idx],
                    'chunk_index': chunk_idx
                }
            })
        
        # Batch upsert (Pinecone recommends batches of 100)
        batch_size = 100
        for i in range(0, len(vectors_to_upsert), batch_size):
            batch = vectors_to_upsert[i:i + batch_size]
            index.upsert(vectors=batch)
            print(f"Upserted batch {i//batch_size + 1}")
            
        print(f"Successfully upserted {len(vectors_to_upsert)} vectors to Pinecone")
        
    except Exception as e:
        print(f"Error upserting to Pinecone: {e}")

# Function to query the index
def query_index(query_text, top_k=5):
    try:
        # First, get embedding for the query
        payload = {
            "model": "models/embedding-001",
            "content": {
                "parts": [{"text": query_text}]
            }
        }
        
        headers = {
            "x-goog-api-key": API_KEY,
            "Content-Type": "application/json"
        }
        
        response = requests.post(ENDPOINT, headers=headers, data=json.dumps(payload))
        
        if response.status_code == 200:
            result = response.json()
            query_embedding = result['embedding']['values']
            
            # Query Pinecone
            query_result = index.query(
                vector=query_embedding,
                top_k=top_k,
                include_metadata=True
            )
            
            return query_result
        else:
            print(f"Error generating query embedding: {response.status_code}, {response.text}")
            return None
            
    except Exception as e:
        print(f"Error querying index: {e}")
        return None

# Full pipeline
if __name__ == "__main__":
    file_path = "HR Policy_CB.pdf"  # Example file path
    
    # Check if file exists
    if not os.path.exists(file_path):
        print(f"File {file_path} not found. Please check the path.")
    else:
        print("Starting document processing pipeline...")
        
        # Extract text
        print("Extracting text from document...")
        text = extract_text(file_path)
        
        if "Error" in text:
            print(f"Text extraction failed: {text}")
        else:
            print(f"Extracted {len(text)} characters")
            
            # Split text into chunks
            print("Chunking text...")
            chunks = chunk_text(text)
            print(f"Created {len(chunks)} chunks")
            
            # Generate embeddings
            print("Generating embeddings...")
            embeddings_with_indices = generate_embeddings_gemini(chunks)
            print(f"Successfully generated {len(embeddings_with_indices)} embeddings")
            
            if embeddings_with_indices:
                # Push to Pinecone
                print("Pushing embeddings to Pinecone...")
                push_embeddings_to_pinecone(embeddings_with_indices, chunks)
                
                # Test query
                print("\nTesting query...")
                query_result = query_index("vacation policy", top_k=3)
                if query_result:
                    print("Query results:")
                    for match in query_result['matches']:
                        print(f"Score: {match['score']:.3f}")
                        print(f"Text: {match['metadata']['text'][:200]}...")
                        print("-" * 50)
            else:
                print("No valid embeddings generated")