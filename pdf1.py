import os
import pytesseract
from PIL import Image
import PyPDF2
import docx
import pdfplumber
import openpyxl  # For handling .xlsx files
import xlrd  # For handling .xls files
import csv  # For handling CSV files

# Set path for Tesseract OCR
pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'

# PDF to text extraction
def pdf_to_text(pdf_path):
    try:
        text = ""
        with open(pdf_path, "rb") as file:
            pdf_reader = PyPDF2.PdfReader(file)
            num_pages = len(pdf_reader.pages)
            for page_num in range(num_pages):
                page = pdf_reader.pages[page_num]
                text += page.extract_text()
        return text
    except Exception as e:
        return f"Error reading PDF: {e}"

# Image to text extraction using Tesseract OCR
def image_to_text(image_path):
    try:
        text = pytesseract.image_to_string(Image.open(image_path))
        return text
    except Exception as e:
        return f"Error reading image: {e}"

# DOCX to text extraction
def docs_to_text(docs_path):
    try:
        doc = docx.Document(docs_path)
        text = ""
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        return text
    except Exception as e:
        return f"Error reading DOCX: {e}"

# TXT to text extraction
def txt_to_text(txt_path):
    try:
        with open(txt_path, 'r', encoding='utf-8') as file:
            return file.read()
    except Exception as e:
        return f"Error reading TXT file: {e}"

# Excel (XLSX) to text extraction (using openpyxl)
def excel_to_text(excel_path):
    try:
        wb = openpyxl.load_workbook(excel_path)
        sheet = wb.active
        text = ""
        for row in sheet.iter_rows(values_only=True):
            text += "\t".join([str(cell) for cell in row]) + "\n"
        return text
    except Exception as e:
        return f"Error reading Excel file (.xlsx): {e}"

# Excel (XLS) to text extraction (using xlrd)
def xls_to_text(xls_path):
    try:
        wb = xlrd.open_workbook(xls_path)
        sheet = wb.sheet_by_index(0)
        text = ""
        for row_num in range(sheet.nrows):
            row = sheet.row(row_num)
            text += "\t".join([str(cell.value) for cell in row]) + "\n"
        return text
    except Exception as e:
        return f"Error reading Excel file (.xls): {e}"

# CSV to text extraction
def csv_to_text(csv_path):
    try:
        with open(csv_path, newline='', encoding='utf-8') as file:
            reader = csv.reader(file)
            text = ""
            for row in reader:
                text += "\t".join(row) + "\n"
        return text
    except Exception as e:
        return f"Error reading CSV file: {e}"

# Main function to extract text based on file extension
def extract_text(file_path):
    _, ext = os.path.splitext(file_path)
    ext = ext.lower()

    if ext == ".pdf":
        return pdf_to_text(file_path)
    elif ext in (".jpg", ".jpeg", ".png", ".bmp", ".gif"):
        return image_to_text(file_path)
    elif ext == ".docx":
        return docs_to_text(file_path)
    elif ext == ".txt":
        return txt_to_text(file_path)
    elif ext == ".xlsx":
        return excel_to_text(file_path)
    elif ext == ".xls":
        return xls_to_text(file_path)
    elif ext == ".csv":
        return csv_to_text(file_path)
    else:
        return "Unsupported file type"

# Example usage:
# Modify file path here to test different files
# file_path = "resume.docx"  # Can be .txt, .xlsx, .xls, .pdf, .docx, etc.
# file_path = "download.jpg"  # Can be .txt, .xlsx, .xls, .pdf, .docx, etc.
file_path = "HR Policy_CB.pdf"  # Can be .txt, .xlsx, .xls, .pdf, .docx, etc.
text = extract_text(file_path)
text_list = text.split("\n")
print(text_list)
